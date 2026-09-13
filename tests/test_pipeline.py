from docx import Document

from docserver import cli


def _criar_docx(caminho, titulo, secao, corpo):
    documento = Document()
    documento.add_heading(titulo, level=1)
    documento.add_heading(secao, level=2)
    documento.add_paragraph(corpo)
    documento.save(caminho)


def test_ingestao_ponta_a_ponta_tres_formatos_busca_encontra_conteudo_dos_tres(tmp_path):
    import pymupdf

    docs_fonte = tmp_path / "docs-fonte"
    docs_normalizado = tmp_path / "docs-normalizado"
    docs_fonte.mkdir()
    docs_normalizado.mkdir()

    (docs_fonte / "markdown.md").write_text(
        "# Documento Markdown\n\n## Seção Markdown\n\n"
        "Este trecho fala sobre renovação de contrato e prazos de vigência.\n",
        encoding="utf-8",
    )

    _criar_docx(
        docs_fonte / "documento.docx",
        "Documento Word",
        "Seção Word",
        "Este trecho fala sobre política de reembolso de despesas de viagem.",
    )

    pdf = pymupdf.open()
    pagina = pdf.new_page()
    pagina.insert_text((72, 72), "Este trecho fala sobre integração com o provedor de pagamentos.")
    pdf.save(docs_fonte / "documento.pdf")
    pdf.close()

    caminho_indice = str(tmp_path / "indice.db")
    relatorio = cli.executar_ingestao(
        docs_fonte, docs_normalizado, caminho_indice, sem_embeddings=True
    )

    assert relatorio["processados"] == 3

    assert len(cli.executar_busca(caminho_indice, "renovação de contrato")) >= 1
    assert len(cli.executar_busca(caminho_indice, "reembolso de despesas")) >= 1
    assert len(cli.executar_busca(caminho_indice, "provedor de pagamentos")) >= 1


def test_arquivo_problematico_no_meio_do_lote_nao_impede_os_demais(tmp_path):
    docs_fonte = tmp_path / "docs-fonte"
    docs_normalizado = tmp_path / "docs-normalizado"
    docs_fonte.mkdir()
    docs_normalizado.mkdir()

    (docs_fonte / "a-antes.md").write_text(
        "# Antes\n\n## Seção\n\nConteúdo do documento que vem antes do problemático.\n",
        encoding="utf-8",
    )
    (docs_fonte / "b-problematico.docx").write_bytes(
        bytes([0x50, 0x4B, 0x03, 0x04]) + bytes(range(256)) * 4
    )
    (docs_fonte / "c-depois.md").write_text(
        "# Depois\n\n## Seção\n\nConteúdo do documento que vem depois do problemático.\n",
        encoding="utf-8",
    )

    caminho_indice = str(tmp_path / "indice.db")
    relatorio = cli.executar_ingestao(
        docs_fonte, docs_normalizado, caminho_indice, sem_embeddings=True
    )

    assert relatorio["processados"] == 2
    assert len(relatorio["falhas"]) == 1
    assert len(cli.executar_busca(caminho_indice, "antes do problemático")) >= 1
    assert len(cli.executar_busca(caminho_indice, "depois do problemático")) >= 1


def test_relatorio_final_traz_as_contagens_corretas(tmp_path):
    docs_fonte = tmp_path / "docs-fonte"
    docs_normalizado = tmp_path / "docs-normalizado"
    docs_fonte.mkdir()
    docs_normalizado.mkdir()

    (docs_fonte / "doc1.md").write_text(
        "# Doc1\n\n## Seção A\n\nConteúdo da seção A com texto suficiente.\n\n"
        "## Seção B\n\nConteúdo da seção B com texto suficiente.\n",
        encoding="utf-8",
    )
    (docs_fonte / "nao-suportado.xyz").write_text("formato desconhecido", encoding="utf-8")

    caminho_indice = str(tmp_path / "indice.db")
    relatorio = cli.executar_ingestao(
        docs_fonte, docs_normalizado, caminho_indice, sem_embeddings=True
    )

    assert relatorio["processados"] == 1
    assert relatorio["chunks"] == 2
    assert len(relatorio["ignorados"]) == 1
    assert len(relatorio["falhas"]) == 0
