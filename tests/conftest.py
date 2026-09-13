import pytest


def _escrever_normalizado(tmp_path, corpo, origem="docs-fonte/exemplo.md", nome="exemplo.md"):
    caminho = tmp_path / nome
    front_matter = (
        "---\n"
        f"origem: {origem}\n"
        "extrator: _extrair_texto_puro\n"
        "ingerido_em: 2026-09-12T10:00:00\n"
        "---\n\n"
    )
    caminho.write_text(front_matter + corpo, encoding="utf-8")
    return caminho


@pytest.fixture
def criar_normalizado(tmp_path):
    def _criar(corpo, origem="docs-fonte/exemplo.md", nome="exemplo.md"):
        return _escrever_normalizado(tmp_path, corpo, origem=origem, nome=nome)

    return _criar


@pytest.fixture
def com_cabecalhos_md(criar_normalizado):
    corpo = (
        "# Manual do Sistema\n\n"
        "## Instalação\n\n"
        "Para instalar, rode o comando de setup e aguarde a conclusão do processo completo.\n\n"
        "## Configuração\n\n"
        "As variáveis de ambiente controlam o comportamento do sistema em produção.\n\n"
        "## Solução de problemas\n\n"
        "Se algo falhar, consulte os logs em /var/log/sistema para mais detalhes.\n"
    )
    return criar_normalizado(corpo, origem="docs-fonte/manual.md", nome="manual.md")


@pytest.fixture
def sem_cabecalhos_md(criar_normalizado):
    paragrafo = (
        "Este é o parágrafo número {n} de exemplo com bastante conteúdo textual "
        "para simular um documento longo sem nenhum cabeçalho de seção."
    )
    paragrafos = [paragrafo.format(n=i) for i in range(80)]
    corpo = "# Documento Sem Secoes\n\n" + "\n\n".join(paragrafos)
    return criar_normalizado(corpo, origem="docs-fonte/corrido.md", nome="corrido.md")


@pytest.fixture
def secao_gigante_md(criar_normalizado):
    paragrafo = (
        "Parágrafo número {n} com texto para forçar a subdivisão da seção gigante em blocos menores."
    )
    paragrafos = [paragrafo.format(n=i) for i in range(120)]
    corpo = (
        "# Documento Com Secao Gigante\n\n"
        "## Introdução\n\nTexto curto de introdução.\n\n"
        "## Referência completa\n\n" + "\n\n".join(paragrafos) + "\n\n"
        "## Conclusão\n\nTexto curto de conclusão.\n"
    )
    return criar_normalizado(corpo, origem="docs-fonte/referencia.md", nome="referencia.md")


@pytest.fixture
def exemplo_docx(tmp_path):
    from docx import Document

    documento = Document()
    documento.add_heading("Título do Documento", level=1)
    documento.add_heading("Primeira Seção", level=2)
    documento.add_paragraph("Parágrafo de exemplo dentro da primeira seção.")

    caminho = tmp_path / "exemplo.docx"
    documento.save(caminho)
    return caminho


@pytest.fixture
def corrompido_docx(tmp_path):
    caminho = tmp_path / "corrompido.docx"
    # assinatura de arquivo ZIP seguida de lixo: força o conversor de docx a ser
    # tentado (e falhar), em vez de cair no fallback de texto puro do markitdown.
    caminho.write_bytes(bytes([0x50, 0x4B, 0x03, 0x04]) + bytes(range(256)) * 4)
    return caminho


@pytest.fixture
def exemplo_pdf(tmp_path):
    import pymupdf

    documento = pymupdf.open()
    pagina = documento.new_page()
    pagina.insert_text((72, 72), "Título do PDF de exemplo com texto extraível.")
    caminho = tmp_path / "exemplo.pdf"
    documento.save(caminho)
    documento.close()
    return caminho


@pytest.fixture
def vazio_pdf(tmp_path):
    import pymupdf

    documento = pymupdf.open()
    documento.new_page()
    caminho = tmp_path / "vazio.pdf"
    documento.save(caminho)
    documento.close()
    return caminho


@pytest.fixture
def docs_fonte(tmp_path):
    caminho = tmp_path / "docs-fonte"
    caminho.mkdir()
    return caminho


@pytest.fixture
def docs_normalizado(tmp_path):
    caminho = tmp_path / "docs-normalizado"
    caminho.mkdir()
    return caminho
